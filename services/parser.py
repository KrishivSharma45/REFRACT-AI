from pathlib import Path


def extract_text_from_txt(file_path):
    """Extract text from a TXT file."""

    path = Path(file_path)

    return path.read_text(
        encoding="utf-8",
        errors="ignore",
    )


def extract_text_from_pdf(file_path):
    """Extract text from a PDF file."""

    try:
        from pypdf import PdfReader
    except ImportError:
        raise RuntimeError(
            "pypdf is not installed. Run: pip install pypdf"
        )

    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    return "\n".join(pages)


def extract_text_from_docx(file_path):
    """Extract text from a DOCX file."""

    try:
        from docx import Document
    except ImportError:
        raise RuntimeError(
            "python-docx is not installed. Run: "
            "pip install python-docx"
        )

    document = Document(file_path)

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            paragraphs.append(paragraph.text)

    # Also extract text from tables.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def extract_text(file_path):
    """
    Automatically extract text based on file extension.

    Supported:
    - PDF
    - DOCX
    - TXT
    """

    if not file_path:
        raise ValueError("No file was provided.")

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    extension = path.suffix.lower()

    if extension == ".txt":
        text = extract_text_from_txt(file_path)

    elif extension == ".pdf":
        text = extract_text_from_pdf(file_path)

    elif extension == ".docx":
        text = extract_text_from_docx(file_path)

    else:
        raise ValueError(
            f"Unsupported file type: {extension}. "
            "Use PDF, DOCX, or TXT."
        )

    text = text.strip()

    if not text:
        raise ValueError(
            f"No readable text was found in {path.name}."
        )

    return text